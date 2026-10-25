#!/usr/bin/env python3
"""
Parse Tokopedia transaction history PDF (Laporan Data Pengguna) into a
categorized Excel spreadsheet.

One product = one row. Multi-item orders are expanded. Failed orders
(Status "Tidak Berhasil") are included.

Usage:
  python parse_tokopedia_trx.py path/to/TrxTokopedia.pdf
  python parse_tokopedia_trx.py path/to/TrxTokopedia.pdf -o output.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypdf import PdfReader

MONTHS = {
    "Jan": 1,
    "Feb": 2,
    "Mar": 3,
    "Apr": 4,
    "Mei": 5,
    "Jun": 6,
    "Jul": 7,
    "Agu": 8,
    "Sep": 9,
    "Okt": 10,
    "Nov": 11,
    "Des": 12,
}

# Order matters: first match wins. Fashion before Pulsa so size "XL" is not
# mistaken for XL Axiata.
CAT_RULES: list[tuple[str, list[str]]] = [
    (
        "Fashion / Pakaian",
        [
            r"\b(celana|kaos kaki|kaos\b|baju|kemeja|rok\b|dress|jaket|hoodie|"
            r"sepatu|sandal|sendal|socks|underwear|boxer|bra\b|hijab|topi|"
            r"dompet|ikat pinggang|fashion|sweater|legging|jeans|blouse|"
            r"outer|cardigan|piyama|pajama)\b"
        ],
    ),
    (
        "Pulsa / Data",
        [
            r"\b(indosat|telkomsel|smartfren|axis|paket data|freedom internet|"
            r"kuota|pulsa|im3|simpati|by\.?u|live\.?on)\b",
            r"\bkartu\s*(xl|tri)\b",
            r"\bxl\s*(pulsa|data|axi|center)\b",
            r"\btri\s*(pulsa|data)\b",
            r"\bpaket\s+merchant\b",
        ],
    ),
    (
        "Tagihan / Token",
        [
            r"\b(token listrik|token pln|\bpln\b|pdam|bpjs|tagihan|pascabayar|"
            r"kota tangerang|tirta benteng)\b"
        ],
    ),
    (
        "Transportasi / Parkir",
        [
            r"\b(parkir|toll|\btol\b|grab|gojek|e-?money|"
            r"flazz|brizzi|jakcard)\b"
        ],
    ),
    (
        "Kesehatan / Obat",
        [
            r"\b(obat|apotek|vitamin|tablet|kapsul|sirup|betadine|paracetamol|"
            r"vermoran|medical|klinik|thermometer|masker medis)\b"
        ],
    ),
    (
        "Anak / Bayi",
        [
            r"\b(bayi|baby|anak|popok|diaper|susu formula|botol susu|mpasi|"
            r"momy|mommy|balita)\b"
        ],
    ),
    (
        "Perawatan Tubuh",
        [
            r"\b(sabun|shampoo|shampo|lulur|scrub|pelembab|moisturizer|serum|"
            r"cream|krim\b|skincare|body lotion|deodorant|pasta gigi|"
            r"sikat gigi|parfum|perfume|cosmetic|makeup|lipstik|bedak|toner|"
            r"facial|sunscreen|sunblock)\b"
        ],
    ),
    (
        "Elektronik / Aksesoris",
        [
            r"\b(charger|kabel data|cable|\busb\b|type-?c|earphone|headphone|"
            r"speaker|powerbank|baterai|battery|adaptor|adapter|keyboard|"
            r"mouse|laptop|iphone|samsung|casing|\bcase\b|\botg\b|\bssd\b|"
            r"\bhdd\b|flashdisk|monitor|\btv\b|lampu led|drawing board|"
            r"tablet|powerbatt|smartwatch|webcam)\b"
        ],
    ),
    (
        "Otomotif",
        [
            r"\b(motor|mobil|vario|\bbeat\b|\bmio\b|helm|oli mesin|\bban\b|"
            r"\baki\b|spion|plat nomor|sarung motor|cover motor|wiper|kampas)\b"
        ],
    ),
    (
        "Sembako / Makanan",
        [
            r"\b(beras|minyak goreng|gula pasir|telur|tepung|\bmie\b|snack|"
            r"makanan|minuman|kopi|\bteh\b|bumbu|sembako|biskuit|coklat|"
            r"keripik)\b"
        ],
    ),
    (
        "Rumah Tangga",
        [
            r"\b(kasur|bantal|selimut|panci|wajan|piring|gelas|sumpit|pisau|"
            r"saringan|ember|sapu|\bpel\b|gantungan|\brak\b|kontainer|wadah|"
            r"grill|panggang|tungku|dapur|keset|gorden|hanger)\b"
        ],
    ),
    (
        "Olahraga",
        [r"\b(olahraga|sport|fitness|yoga|dumbbell|sepeda|\bbola\b|running|\bgym\b)\b"],
    ),
]

TXN_PATTERN = re.compile(
    r"(?:No\.\s*Invoice\n([^\n]+)\n)?"
    r"Status\n(Berhasil|Tidak Berhasil)\n"
    r"Tanggal Order\n([^\n]+)\n"
    r"Nama Produk\n(.*?)"
    r"Jumlah Total\n([^\n]+)"
    r"(?:\nNama Penjual\n([^\n]+))?",
    re.S,
)

HEADERS = [
    "No",
    "Tanggal",
    "Kategori",
    "Nama Barang",
    "Nilai (Rp)",
    "No. Invoice",
    "Penjual",
    "Status",
]

ADJUSTMENT_NAME = "Potongan harga (tidak teralokasi)"
ADJUSTMENT_CATEGORY = "Penyesuaian"

RP_NUMBER_FORMAT = "#,##0"

# Suffix-only: a leading (.*?) next to \s* backtracks super-linearly.
_QTY_PRICE_RE = re.compile(r"\(x(\d+)\)\s*@Rp([\d.]+)\s*$")
_PRICE_RE = re.compile(r"@Rp([\d.]+)\s*$")
_QTY_RE = re.compile(r"\(x(\d+)\)\s*$")


def extract_text(pdf_path: Path) -> str:
    reader = PdfReader(str(pdf_path))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    text = "\n".join(parts)
    text = re.sub(r"\n?Confidential - Downloaded on [^\n]+\n?", "\n", text)
    idx = text.find("Status\n")
    if idx < 0:
        raise ValueError("Tidak menemukan blok Status di PDF. Cek format export.")
    return text[idx:]


def parse_date(raw: str) -> str:
    """Convert 'Jumat, 29 Des 2023 15:58:26' -> '2023-12-29'."""
    m = re.search(r"(\d{1,2})\s+(\w+)\s+(\d{4})", raw)
    if not m:
        return raw.strip()
    day, mon, year = int(m.group(1)), m.group(2), int(m.group(3))
    month = MONTHS.get(mon)
    if not month:
        return raw.strip()
    return f"{year:04d}-{month:02d}-{day:02d}"


def parse_rp(raw: str) -> int:
    digits = re.sub(r"[^\d]", "", raw.replace("Rp", ""))
    return int(digits) if digits else 0


def categorize(name: str) -> str:
    low = name.lower()
    for cat, patterns in CAT_RULES:
        for pat in patterns:
            if re.search(pat, low, re.I):
                return cat
    return "Lainnya"


def _one_item(rest: str) -> tuple[str, int | None]:
    m = _QTY_PRICE_RE.search(rest)
    if m:
        name = rest[: m.start()].strip()
        qty, unit = int(m.group(1)), parse_rp(m.group(2))
        return name + f" (x{qty})", unit * qty

    m = _PRICE_RE.search(rest)
    if m:
        name = rest[: m.start()].strip()
        qm = _QTY_RE.search(name)
        if qm:
            qty = int(qm.group(1))
            name_clean = name[: qm.start()].strip()
            unit = parse_rp(m.group(1))
            return name_clean + f" (x{qty})", unit * qty
        return name, parse_rp(m.group(1))

    m = _QTY_RE.search(rest)
    if m:
        return rest[: m.start()].strip() + f" (x{m.group(1)})", None

    return rest, None


def _indexed_item_body(line: str) -> str | None:
    """Text after a leading '12. ' index. None when the line is not indexed."""
    m = re.match(r"\d+\.\s", line)
    if m is None:
        return None
    return line[m.end() :].strip()


def parse_products(prod_text: str) -> list[tuple[str, int | None]]:
    """Return list of (product_name, value_or_None)."""
    lines = [ln.strip() for ln in prod_text.strip().splitlines() if ln.strip()]
    joined: list[str] = []
    for ln in lines:
        if re.match(r"^\d+\.\s", ln):
            joined.append(ln)
        elif re.match(r"^@Rp", ln) and joined:
            joined[-1] = joined[-1] + " " + ln
        elif joined:
            joined[-1] = joined[-1] + " " + ln
        else:
            joined.append(ln)

    if any(re.match(r"^\d+\.\s", j) for j in joined):
        items: list[tuple[str, int | None]] = []
        for j in joined:
            body = _indexed_item_body(j)
            if body is not None:
                items.append(_one_item(body))
        return items

    full = re.sub(r"\s+", " ", " ".join(joined)).strip()
    return [_one_item(full)]


def _resolve_item_values(
    products: list[tuple[str, int | None]], total: int
) -> list[tuple[str, int, str | None]]:
    """Return (name, value, category override).

    One item with @Rp: Jumlah Total wins when it is lower (discount on that
    item). A higher total is shipping or a fee, so the catalog price stays.
    One item without @Rp uses Jumlah Total.

    Several items keep @Rp × qty. A lower Jumlah Total is one negative row
    after those items; which item was discounted is unknown. A higher total
    is left off the sheet.
    """
    if len(products) == 1:
        name, value = products[0]
        if value is None or total < value:
            return [(name, total, None)]
        return [(name, value, None)]
    if not products:
        return []

    rows = [(name, 0 if value is None else value, None) for name, value in products]
    if any(value is None for _, value in products):
        return rows
    gap = total - sum(value for _, value, _ in rows)
    if gap < 0:
        rows.append((ADJUSTMENT_NAME, gap, ADJUSTMENT_CATEGORY))
    return rows


def parse_transactions(text: str) -> tuple[list[dict], dict]:
    """
    Parse all orders into item rows.

    Invoice rule: No. Invoice immediately *before* Status belongs to that order.
    (Trailing invoice after Nama Penjual belongs to the next order.)
    """
    matches = list(TXN_PATTERN.finditer(text))
    rows: list[dict] = []
    status_counter: Counter[str] = Counter()
    no_invoice = 0

    for m in matches:
        invoice, status, tanggal_raw, prod_text, total_raw, seller = m.groups()
        invoice = (invoice or "").strip()
        if not invoice:
            no_invoice += 1
        seller = (seller or "").strip()
        tanggal = parse_date(tanggal_raw)
        total = parse_rp(total_raw)
        status_counter[status] += 1

        for name, nilai, kategori in _resolve_item_values(parse_products(prod_text), total):
            rows.append(
                {
                    "tanggal": tanggal,
                    "kategori": kategori or categorize(name),
                    "nama": name,
                    "nilai": nilai,
                    "invoice": invoice,
                    "penjual": seller,
                    "status": status,
                }
            )

    meta = {
        "orders": len(matches),
        "items": len(rows),
        "status": dict(status_counter),
        "no_invoice": no_invoice,
    }
    return rows, meta


def _write_item_rows(ws, rows: list[dict], body_font: Font, thin: Border, fail_fill: PatternFill) -> None:
    for idx, r in enumerate(rows, start=1):
        ws.append(
            [
                idx,
                r["tanggal"],
                r["kategori"],
                r["nama"],
                r["nilai"],
                r["invoice"],
                r["penjual"],
                r["status"],
            ]
        )
        row_i = idx + 1
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row_i, col)
            cell.font = body_font
            cell.border = thin
            if r["status"] != "Berhasil":
                cell.fill = fail_fill
        ws.cell(row_i, 1).alignment = Alignment(horizontal="center")
        ws.cell(row_i, 5).number_format = RP_NUMBER_FORMAT


def _write_summary(ws2, rows: list[dict], meta: dict, body_font: Font) -> None:
    ws2["A1"] = "Ringkasan Trx Tokopedia"
    ws2["A1"].font = Font(name="Arial", bold=True, size=12)
    ws2["A3"] = "Total baris item"
    ws2["B3"] = meta["items"]
    ws2["A4"] = "Total transaksi (order)"
    ws2["B4"] = meta["orders"]
    ws2["A5"] = "Berhasil"
    ws2["B5"] = meta["status"].get("Berhasil", 0)
    ws2["A6"] = "Tidak Berhasil"
    ws2["B6"] = meta["status"].get("Tidak Berhasil", 0)
    ws2["A7"] = "Tanpa No. Invoice"
    ws2["B7"] = meta["no_invoice"]

    for cell in ("A3", "A4", "A5", "A6", "A7"):
        ws2[cell].font = Font(name="Arial", bold=True, size=10)
    for cell in ("B3", "B4", "B5", "B6", "B7"):
        ws2[cell].font = body_font

    cat_sum: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    cat_ok: dict[str, int] = defaultdict(int)
    for r in rows:
        cat_sum[r["kategori"]][0] += 1
        cat_sum[r["kategori"]][1] += r["nilai"]
        if r["status"] == "Berhasil":
            cat_ok[r["kategori"]] += r["nilai"]

    ws2["A9"] = "Per kategori (semua status)"
    ws2["A9"].font = Font(name="Arial", bold=True)
    ws2["A10"] = "Kategori"
    ws2["B10"] = "Jumlah baris"
    ws2["C10"] = "Total Nilai (Rp)"
    for c in ("A10", "B10", "C10"):
        ws2[c].font = Font(name="Arial", bold=True, size=10)

    for i, (cat, (n, s)) in enumerate(
        sorted(cat_sum.items(), key=lambda x: -x[1][1]), start=11
    ):
        ws2[f"A{i}"] = cat
        ws2[f"B{i}"] = n
        ws2[f"C{i}"] = s
        ws2[f"C{i}"].number_format = RP_NUMBER_FORMAT
        for col in "ABC":
            ws2[f"{col}{i}"].font = body_font

    ws2["E9"] = "Total nilai Status=Berhasil"
    ws2["E9"].font = Font(name="Arial", bold=True)
    ws2["E10"] = "Kategori"
    ws2["F10"] = "Total (Rp)"
    ws2["E10"].font = Font(name="Arial", bold=True, size=10)
    ws2["F10"].font = Font(name="Arial", bold=True, size=10)
    for i, (cat, s) in enumerate(sorted(cat_ok.items(), key=lambda x: -x[1]), start=11):
        ws2[f"E{i}"] = cat
        ws2[f"F{i}"] = s
        ws2[f"F{i}"].number_format = RP_NUMBER_FORMAT
        ws2[f"E{i}"].font = body_font
        ws2[f"F{i}"].font = body_font

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 18
    ws2.column_dimensions["E"].width = 28
    ws2.column_dimensions["F"].width = 16


def write_xlsx(rows: list[dict], meta: dict, out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Trx Tokopedia"

    header_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    header_fill = PatternFill("solid", fgColor="333333")
    fail_fill = PatternFill("solid", fgColor="FFE8E8")
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    ws.append(HEADERS)
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin

    _write_item_rows(ws, rows, body_font, thin, fail_fill)

    widths = {
        "A": 6,
        "B": 12,
        "C": 22,
        "D": 70,
        "E": 12,
        "F": 30,
        "G": 24,
        "H": 14,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A1:H{len(rows) + 1}"
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Ringkasan")
    _write_summary(ws2, rows, meta, body_font)

    # Catatan
    ws3 = wb.create_sheet("Catatan")
    ws3["A1"] = "Catatan"
    ws3["A1"].font = Font(name="Arial", bold=True, size=11)
    notes = [
        "Sumber: export PDF Riwayat Transaksi Tokopedia (Laporan Data Pengguna).",
        "1 barang dengan @Rp: Nilai = Jumlah Total bila lebih kecil (diskon). Bila lebih besar, tetap @Rp.",
        "1 barang tanpa @Rp: Nilai = Jumlah Total.",
        "Multi-item: tiap barang @Rp × qty. Bila Jumlah Total lebih kecil, baris berikutnya "
        "adalah Potongan harga (tidak teralokasi), kategori Penyesuaian.",
        "Selisih positif (ongkir / biaya layanan) tidak dicantumkan.",
        "No. Invoice = baris invoice tepat sebelum Status order.",
        "Baris Tidak Berhasil di-highlight merah muda.",
        "Kategori otomatis (kata kunci); koreksi manual di kolom Kategori bila perlu.",
        f"Order ter-parse: {meta['orders']} | Item rows: {meta['items']} | "
        f"Tanpa invoice: {meta['no_invoice']}.",
    ]
    for i, note in enumerate(notes, start=3):
        ws3[f"A{i}"] = note
        ws3[f"A{i}"].font = Font(name="Arial", size=10)
    ws3.column_dimensions["A"].width = 130

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Parse Tokopedia transaction PDF into categorized Excel."
    )
    parser.add_argument("pdf", type=Path, help="Path ke PDF export Tokopedia")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Path output .xlsx (default: <pdf_stem>_kategori.xlsx)",
    )
    args = parser.parse_args(argv)

    pdf_path: Path = args.pdf
    if not pdf_path.is_file():
        print(f"File tidak ditemukan: {pdf_path}", file=sys.stderr)
        return 1

    out_path = args.output or pdf_path.with_name(f"{pdf_path.stem}_kategori.xlsx")

    print(f"Membaca: {pdf_path}")
    text = extract_text(pdf_path)
    rows, meta = parse_transactions(text)
    write_xlsx(rows, meta, out_path)

    total_ok = sum(r["nilai"] for r in rows if r["status"] == "Berhasil")
    print(f"Order: {meta['orders']} | Item: {meta['items']}")
    print(f"Status: {meta['status']}")
    print(f"Tanpa invoice: {meta['no_invoice']}")
    print(f"Total nilai Berhasil: Rp{total_ok:,}")
    print(f"Tersimpan: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
